from pathlib import Path
from math import atan2, cos, sin, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\table\jiuyuan")
TEMPLATE = Path(r"D:\table\4928bfc9-c8a1-484f-848e-9dcb9f15ba7a(7).docx")
OUT = ROOT / "星闪灾后救援系统设计报告.docx"
ASSET_DIR = ROOT / "report_assets"


BLUE = "#1F4E79"
LIGHT_BLUE = "#D9EAF7"
PALE_BLUE = "#EEF5FB"
GRAY = "#F2F4F7"
DARK = "#1F2933"
GREEN = "#2F855A"
ORANGE = "#C05621"
RED = "#C53030"


def get_font(size=32, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc") if bold else Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def wrap_text(draw, text, font, width):
    lines = []
    current = ""
    for ch in text:
        test = current + ch
        if draw.textlength(test, font=font) <= width or not current:
            current = test
        else:
            lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def draw_centered(draw, box, text, font, fill=DARK, spacing=6):
    x1, y1, x2, y2 = box
    lines = []
    for part in text.split("\n"):
        lines.extend(wrap_text(draw, part, font, x2 - x1 - 28))
    line_heights = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        line_heights.append(bbox[3] - bbox[1])
    total_h = sum(line_heights) + spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, h in zip(lines, line_heights):
        w = draw.textlength(line, font=font)
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=font, fill=fill)
        y += h + spacing


def draw_round_box(draw, box, title, body, fill, outline=BLUE):
    title_font = get_font(30, bold=True)
    body_font = get_font(23)
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=3)
    x1, y1, x2, y2 = box
    draw_centered(draw, (x1 + 10, y1 + 14, x2 - 10, y1 + 62), title, title_font, fill=outline)
    body_lines = []
    for part in body.split("\n"):
        body_lines.extend(wrap_text(draw, part, body_font, x2 - x1 - 36))
    y = y1 + 82
    for line in body_lines:
        draw.text((x1 + 22, y), line, font=body_font, fill=DARK)
        y += 31


def arrow(draw, start, end, fill=BLUE, width=5):
    draw.line([start, end], fill=fill, width=width)
    angle = atan2(end[1] - start[1], end[0] - start[0])
    head = 20
    points = [
        end,
        (end[0] - head * cos(angle - pi / 6), end[1] - head * sin(angle - pi / 6)),
        (end[0] - head * cos(angle + pi / 6), end[1] - head * sin(angle + pi / 6)),
    ]
    draw.polygon(points, fill=fill)


def save_architecture(path):
    img = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(img)
    title_font = get_font(46, True)
    subtitle_font = get_font(26)
    draw.text((70, 45), "基于星闪互联的灾后救援系统总体架构", font=title_font, fill=BLUE)
    draw.text((72, 108), "从两块 WS63 基础闭环扩展到现场网关与指挥后台", font=subtitle_font, fill="#4A5568")

    boxes = {
        "victim": (80, 230, 420, 430),
        "rescuer": (80, 600, 420, 800),
        "gateway": (680, 420, 1030, 620),
        "backend": (1270, 250, 1660, 430),
        "storage": (1270, 520, 1660, 700),
        "web": (1270, 790, 1660, 940),
    }
    draw_round_box(draw, boxes["victim"], "受灾人员终端", "WS63 SLE Server\n身份广播 / 连接\nSpO2、HR、SBP、电量、SOS", PALE_BLUE)
    draw_round_box(draw, boxes["rescuer"], "救援人员终端", "WS63 SLE Client\n扫描连接 / RSSI\n距离等级、趋势、OLED 显示", PALE_BLUE)
    draw_round_box(draw, boxes["gateway"], "现场星闪网关", "星闪接入汇聚\n本地缓存\nMQTT/HTTPS 转发", "#F7FAFC")
    draw_round_box(draw, boxes["backend"], "Spring Boot 后端", "设备、告警、任务\n优先级计算\n实时接口", "#F7FAFC")
    draw_round_box(draw, boxes["storage"], "MySQL / Redis", "历史数据入库\n实时状态缓存\n优先级队列", "#F7FAFC")
    draw_round_box(draw, boxes["web"], "Vue3 指挥后台", "态势总览\n人员列表\n任务派发与复盘", "#F7FAFC")

    arrow(draw, (420, 330), (680, 470))
    draw.text((482, 350), "SLE 数据上报", font=get_font(22), fill=BLUE)
    arrow(draw, (420, 700), (680, 570))
    draw.text((482, 655), "SLE/有线/4G", font=get_font(22), fill=BLUE)
    arrow(draw, (420, 515), (80, 515))
    draw.text((150, 485), "两板基础版：SLE 连接 + RSSI 粗距离 + 趋势", font=get_font(24, True), fill=GREEN)
    arrow(draw, (1030, 520), (1270, 340))
    draw.text((1080, 405), "MQTT / HTTPS", font=get_font(22), fill=BLUE)
    arrow(draw, (1465, 430), (1465, 520))
    arrow(draw, (1465, 700), (1465, 790))
    draw.text((1490, 735), "WebSocket/SSE", font=get_font(22), fill=BLUE)
    img.save(path)


def save_workflow(path):
    img = Image.new("RGB", (1800, 860), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), "两块 WS63 基础定位工作流程", font=get_font(46, True), fill=BLUE)
    draw.text((72, 105), "当前原型不输出绝对坐标，重点完成发现目标、粗距离估算和靠近/远离趋势判断", font=get_font(25), fill="#4A5568")
    steps = [
        ("1 受灾端广播", "SLE Server 广播\n设备名 sle_uart_server\n等待救援端连接"),
        ("2 救援端扫描", "SLE Client 扫描目标\n发现后建立连接\n读取初始 RSSI"),
        ("3 生命体征上报", "每 1 s 发送 JSON\nid / spo2 / hr / sbp\nbattery / sos"),
        ("4 RSSI 估距", "读取连接 RSSI\n经验模型换算距离\n输出 VERY_NEAR~FAR"),
        ("5 OLED 与串口", "显示 RSSI、DIST\nLEVEL、TREND\nID、SpO2、HR、SOS"),
    ]
    x = 70
    y = 250
    w = 300
    h = 270
    for i, (title, body) in enumerate(steps):
        fill = PALE_BLUE if i % 2 == 0 else "#F7FAFC"
        draw_round_box(draw, (x, y, x + w, y + h), title, body, fill)
        if i < len(steps) - 1:
            arrow(draw, (x + w + 18, y + h / 2), (x + w + 80, y + h / 2), fill=ORANGE, width=5)
        x += w + 95
    draw.rounded_rectangle((160, 635, 1640, 770), radius=18, fill="#FFF7ED", outline=ORANGE, width=3)
    note = "技术边界：两块普通 WS63 板不能天然得到稳定二维坐标或绝对方向；报告将当前实现定义为 RSSI 粗定位与趋势导航，SLP/HADM 或多锚点定位作为增强方案。"
    draw_centered(draw, (190, 645, 1610, 760), note, get_font(28), fill="#7C2D12")
    img.save(path)


def save_oled(path):
    img = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 45), "WS63 OLED 终端界面示意", font=get_font(46, True), fill=BLUE)
    draw.text((72, 110), "根据当前源码中的 SSD1306 显示字段绘制，用于说明现场可视化效果", font=get_font(25), fill="#4A5568")
    mono = get_font(34, True)
    small = get_font(25)

    def device(x, y, title, lines, accent):
        draw.rounded_rectangle((x, y, x + 620, y + 500), radius=42, fill="#1A202C", outline=accent, width=8)
        draw.rounded_rectangle((x + 52, y + 55, x + 568, y + 405), radius=8, fill="#000000", outline="#CBD5E0", width=3)
        yy = y + 75
        for idx, line in enumerate(lines):
            color = "#68D391" if idx == 0 else "#F7FAFC"
            draw.text((x + 82, yy), line, font=mono, fill=color)
            yy += 40
        draw.text((x + 88, y + 435), title, font=small, fill="#E2E8F0")

    device(
        180,
        240,
        "救援端 WS63：RSSI 估距与目标状态",
        [
            "RESCUE TRACK",
            "LINK CONNECTED",
            "RSSI -61",
            "DIST 10.8m",
            "LEVEL NEAR",
            "ID VD-001",
            "SpO2 90 HR 124",
            "SOS Y BAT 80",
        ],
        "#63B3ED",
    )
    device(
        1000,
        240,
        "受灾端 WS63：生命体征模拟上报",
        [
            "RESCUE VICTIM",
            "LINK ADVERTISING",
            "ID VD-001 SOS YES",
            "SEQ 12",
            "SpO2 90 HR 124",
            "SBP 100 BAT 80",
            "SENDING SLE REPORT",
            "NO TEMP/HUMI DEMO",
        ],
        "#68D391",
    )
    img.save(path)


def save_priority(path):
    img = Image.new("RGB", (1800, 900), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 45), "救援优先级计算结构", font=get_font(46, True), fill=BLUE)
    draw.text((72, 110), "后台增强版按生命风险、失联静止、位置可达性、环境风险和 SOS 形成可解释排序", font=get_font(25), fill="#4A5568")
    blocks = [
        ("生命风险 45%", "SpO2、HR、SBP/DBP 趋势\n血压估算需降低权重并标记 estimated_bp", RED),
        ("失联/静止 20%", "跌倒、长时间静止、数据过期、离线", ORANGE),
        ("位置可达 15%", "距离最近救援员、定位质量、未知位置惩罚", BLUE),
        ("环境风险 10%", "高温、可燃气体、缺氧、CO 等安全风险", "#7A5A00"),
        ("SOS 10%", "主动求救信号作为高优先级触发因子", GREEN),
    ]
    y = 205
    for title, body, color in blocks:
        draw.rounded_rectangle((110, y, 750, y + 105), radius=18, fill="#F7FAFC", outline=color, width=4)
        draw.text((135, y + 18), title, font=get_font(30, True), fill=color)
        draw.text((350, y + 18), body, font=get_font(23), fill=DARK)
        arrow(draw, (750, y + 52), (940, 455), fill=color, width=4)
        y += 125
    draw.rounded_rectangle((940, 310, 1280, 600), radius=28, fill=PALE_BLUE, outline=BLUE, width=4)
    draw_centered(draw, (960, 330, 1260, 580), "priority_score\n= 加权综合评分\n输出 RED / ORANGE /\nYELLOW / GREEN / GRAY", get_font(30, True), fill=BLUE)
    draw.rounded_rectangle((1390, 310, 1680, 600), radius=28, fill="#F0FFF4", outline=GREEN, width=4)
    draw_centered(draw, (1410, 330, 1660, 580), "可解释原因\n分数\n等级\n置信度\n派发建议", get_font(31, True), fill=GREEN)
    arrow(draw, (1280, 455), (1390, 455), fill=GREEN, width=5)
    img.save(path)


def set_run_font(run, east="宋体", west="Times New Roman", size=None, bold=None, color=None):
    run.font.name = west
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), west)
    rfonts.set(qn("w:hAnsi"), west)
    rfonts.set(qn("w:eastAsia"), east)


def set_cell_text(cell, text, bold=False, fill=None, align=WD_ALIGN_PARAGRAPH.LEFT, size=Pt(10)):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=DARK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill.lstrip("#"))
        cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(table, top=90, start=120, bottom=90, end=120):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")
    set_cell_margins(table)


def clear_body(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def setup_document(doc):
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.4)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.75)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = "Times New Roman"
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size, color, before, after in [
        ("Heading 1", 15, BLUE, 12, 8),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_paragraph(doc, text="", first_line=True, bold=False):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_run_font(r, size=Pt(10.5), bold=bold, color=DARK)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    size = {1: Pt(15), 2: Pt(13), 3: Pt(12)}[level]
    east = "黑体" if level in (1, 2) else "宋体"
    set_run_font(r, east=east, size=size, bold=True, color=BLUE if level < 3 else "1F4D78")
    return p


def add_bullets(doc, items):
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    nums = numbering.findall(qn("w:num"))
    next_num_id = max([int(n.get(qn("w:numId"))) for n in nums] or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(next_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), "1")
    num.append(abstract)
    numbering.append(num)
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        ppr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id = OxmlElement("w:numId")
        num_id.set(qn("w:val"), str(next_num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id)
        ppr.append(num_pr)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=Pt(10.5), color=DARK)


def add_caption(doc, text, kind="figure"):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4 if kind == "figure" else 8)
    p.paragraph_format.space_after = Pt(6 if kind == "figure" else 4)
    r = p.add_run(text)
    set_run_font(r, size=Pt(10), bold=False, color="4A5568")
    return p


def add_picture(doc, path, caption, width=5.9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    if doc.inline_shapes:
        doc.inline_shapes[-1]._inline.docPr.set("descr", caption)
    add_caption(doc, caption)


def add_table_caption(doc, text):
    add_caption(doc, text, kind="table")


def add_data_table(doc, headers, rows, widths_cm, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths_cm)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, fill=header_fill, align=WD_ALIGN_PARAGRAPH.CENTER)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(str(value)) < 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], str(value), align=align)
    doc.add_paragraph()
    return table


def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [16.0])
    cell = table.cell(0, 0)
    set_cell_text(cell, text, fill="F7FAFC", size=Pt(9))
    for p in cell.paragraphs:
        for run in p.runs:
            set_run_font(run, east="宋体", west="Consolas", size=Pt(9), color="2D3748")
    doc.add_paragraph()


def add_toc(doc):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目  录")
    set_run_font(r, east="黑体", size=Pt(15), bold=True, color=BLUE)
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    run = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    run._r.append(instr)
    run = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "separate")
    run._r.append(fld)
    run = p.add_run("目录正在生成，请在 Word 中更新域。")
    set_run_font(run, size=Pt(10.5), color="718096")
    run = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "end")
    run._r.append(fld)


def build_assets():
    ASSET_DIR.mkdir(exist_ok=True)
    paths = {
        "architecture": ASSET_DIR / "architecture.png",
        "workflow": ASSET_DIR / "workflow.png",
        "oled": ASSET_DIR / "oled.png",
        "priority": ASSET_DIR / "priority.png",
    }
    save_architecture(paths["architecture"])
    save_workflow(paths["workflow"])
    save_oled(paths["oled"])
    save_priority(paths["priority"])
    return paths


def build_report():
    assets = build_assets()
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    setup_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(84)
    r = p.add_run("全国大学生物联网设计竞赛")
    set_run_font(r, east="黑体", size=Pt(22), bold=True, color=DARK)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(80)
    r = p.add_run("基于星闪互联的灾后救援系统")
    set_run_font(r, east="黑体", size=Pt(24), bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(210)
    r = p.add_run("全国大学生物联网设计竞赛组委会")
    set_run_font(r, east="宋体", size=Pt(14), bold=False, color=DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2026年6月")
    set_run_font(r, east="宋体", size=Pt(12), bold=False, color=DARK)
    doc.add_page_break()

    # Team information table
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("参赛作品信息")
    set_run_font(r, east="黑体", size=Pt(15), bold=True, color=BLUE)
    table = doc.add_table(rows=7, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [4.2, 11.8])
    info_rows = [
        ("学校名称：", "填写学校名称"),
        ("团队名称：", "星闪灾后救援系统设计团队"),
        ("", ""),
        ("队长：", "填写姓名"),
        ("队员1：", "填写姓名"),
        ("队员2：", "填写姓名"),
        ("队员3：", "填写姓名"),
    ]
    for row, (k, v) in zip(table.rows, info_rows):
        set_cell_text(row.cells[0], k, bold=True, fill=GRAY, align=WD_ALIGN_PARAGRAPH.RIGHT)
        set_cell_text(row.cells[1], v, align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_page_break()

    # Abstract and TOC
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于星闪互联的灾后救援系统")
    set_run_font(r, east="黑体", size=Pt(18), bold=True, color=BLUE)
    add_heading(doc, "摘要", 1)
    abstract = (
        "灾后救援现场常存在公网中断、建筑遮挡、烟尘噪声、二次灾害风险高等问题，传统依赖人工呼喊、普通短距无线或后台网络的搜救方式难以及时发现被困人员并判断救援优先级。"
        "本作品设计并实现一套基于星闪 NearLink 的灾后救援系统，系统由受灾人员终端、救援人员终端、现场网关和指挥后台组成。当前原型以两块 WS63 开发板完成最小闭环：受灾端作为 SLE Server 广播并周期发送模拟生命体征 JSON 数据，救援端作为 SLE Client 自动扫描名为 sle_uart_server 的目标、建立连接、读取 RSSI、估算粗距离等级，并根据连续距离变化判断 CLOSING、AWAY、STABLE 趋势；两端均接入 SSD1306 OLED，用于显示连接状态、目标编号、RSSI、距离等级、SpO2、HR、SBP、电量和 SOS 状态。"
        "完整系统设计在此基础上扩展现场网关、Spring Boot 后端、MySQL/Redis 数据存储、Vue3 指挥后台和救援优先级算法，实现本地自治搜救、实时态势展示、告警推送和任务派发。报告严格区分基础原型能力与工程增强能力：两块板当前不能直接输出稳定二维坐标或绝对方向，精确定位需通过 SLP/HADM 测距、多锚点部署或 IMU 轨迹融合进一步实现。"
    )
    add_paragraph(doc, abstract)
    p = doc.add_paragraph(style="Normal")
    r = p.add_run("关键词：")
    set_run_font(r, east="黑体", size=Pt(10.5), bold=True, color=DARK)
    r = p.add_run("星闪 NearLink；WS63；SLE 通信；灾后救援；RSSI 粗定位；生命体征监测；OLED 显示；救援优先级")
    set_run_font(r, size=Pt(10.5), color=DARK)
    doc.add_paragraph()
    add_toc(doc)
    doc.add_page_break()

    # 1
    add_heading(doc, "1 设计需求分析", 1)
    add_heading(doc, "1.1 应用场景", 2)
    add_paragraph(
        doc,
        "本作品面向地震、矿难、火灾、坍塌、隧道事故和大型公共场所灾害后的现场搜救。上述场景具有空间结构破坏严重、人员分布未知、通信基础设施不可靠、救援窗口短和救援人员安全风险高等特点。系统设计目标不是替代专业生命探测仪或医疗诊断设备，而是在低成本可穿戴终端和短距无线互联基础上，为救援人员提供“先发现、再接近、再评估、再派发”的辅助决策工具。",
    )
    add_picture(doc, assets["architecture"], "图1-1 系统总体架构图")
    add_heading(doc, "1.2 现有方式的不足", 2)
    add_paragraph(
        doc,
        "传统灾后搜救常依赖人工搜索、对讲机、手机定位或单一生命探测设备。人工搜索效率受噪声、烟尘、遮挡影响大；公网定位在基站损毁或室内遮挡环境下不稳定；普通蓝牙或 Wi-Fi 方案在复杂电磁环境中可靠性不足；单一生命探测设备缺少人员身份、电量、告警和任务闭环信息。救援指挥还需要同时判断生命风险、距离远近、现场危险和救援人员分配状态，单点设备难以支撑全局调度。",
    )
    add_heading(doc, "1.3 需求分解", 2)
    add_table_caption(doc, "表1-1 灾后救援需求与系统响应关系")
    add_data_table(
        doc,
        ["需求类别", "关键问题", "系统响应", "当前实现边界"],
        [
            ["目标发现", "不知道被困人员是否在附近", "受灾终端广播身份，救援端扫描并连接", "已在两块 WS63 上实现 SLE 扫描连接"],
            ["接近判断", "无法判断移动方向是否正确", "连续 RSSI 估距并输出靠近/远离趋势", "基础版为粗距离和趋势，不是绝对方向"],
            ["生命状态", "只知道有人，无法判断紧急程度", "上报 SpO2、HR、SBP、电量、SOS 等字段", "当前为模拟数据，后续接入 PPG 传感器"],
            ["现场调度", "多目标救援优先级难以排序", "后台按生命风险、位置、环境、SOS 计算分级", "属于完整系统设计和增强实现"],
            ["断网可用", "公网或后台不可用时仍需搜救", "终端本地显示距离、趋势、目标生命状态", "已在 OLED 与串口输出中体现"],
        ],
        [2.6, 4.3, 5.0, 4.1],
    )
    add_heading(doc, "1.4 设计约束", 2)
    add_bullets(
        doc,
        [
            "本地自治优先：救援端必须在无后台、无公网条件下完成目标发现、连接、状态显示和接近趋势判断。",
            "技术边界清晰：两块 WS63 基础版不能宣称稳定二维坐标定位，RSSI 估距仅用于粗略距离等级和趋势辅助。",
            "医疗边界清晰：血氧、心率和血压趋势用于救援辅助，正式医疗判断必须由专业设备和医护人员完成。",
            "工程可扩展：基础通信协议、数据字段和界面布局需预留网关、后台、多锚点定位和真实传感器接入空间。",
        ],
    )

    # 2
    add_heading(doc, "2 特色与创新", 1)
    add_heading(doc, "2.1 星闪短距互联用于救援近场发现", 2)
    add_paragraph(
        doc,
        "作品把星闪 NearLink 的短距低时延连接能力用于灾后近场搜救。受灾人员终端主动广播，救援人员终端进入现场后无需依赖公网即可扫描目标并建立 SLE 连接。相较单纯后台定位方案，本设计把关键判断前移到救援人员手持终端，适合断网、遮挡和应急部署环境。",
    )
    add_heading(doc, "2.2 两板基础闭环可平滑升级", 2)
    add_paragraph(
        doc,
        "系统从两块 WS63 开发板起步，基础版完成 SLE 连接、模拟生命体征上报、RSSI 粗距离估计和 OLED 显示；增强版可以把 RSSI 数据源替换为 SLP/HADM 测距；工程版继续加入 IMU 运动轨迹、多锚点和现场网关。这样的分层设计避免把未验证的高精度定位能力写死在系统核心流程中，便于竞赛演示和后续工程扩展。",
    )
    add_picture(doc, assets["workflow"], "图2-1 两块 WS63 基础定位工作流程图")
    add_heading(doc, "2.3 生命体征与位置趋势联合判断", 2)
    add_paragraph(
        doc,
        "普通无线寻物方案通常只提供距离或信号强度，本作品把目标身份、生命体征摘要、电量、SOS 状态与距离趋势合并显示。救援人员不仅能判断“是否靠近目标”，还能判断“目标是否高危、是否主动求救、电量是否即将耗尽”，从而减少盲目搜索和重复确认时间。",
    )
    add_heading(doc, "2.4 可解释救援优先级", 2)
    add_paragraph(
        doc,
        "后台增强版不只输出一个黑箱分数，而是把 SpO2、心率、血压趋势、SOS、失联时间、环境风险和距离可达性拆分为可解释原因。指挥员可以看到红、橙、黄、绿、灰分级及对应依据，必要时进行人工调整，降低算法误判带来的调度风险。",
    )
    add_picture(doc, assets["priority"], "图2-2 救援优先级算法结构图")

    # 3
    add_heading(doc, "3 功能设计", 1)
    add_heading(doc, "3.1 系统角色与终端功能", 2)
    add_table_caption(doc, "表3-1 功能模块设计")
    add_data_table(
        doc,
        ["模块", "主要功能", "输入数据", "输出结果"],
        [
            ["受灾人员终端", "身份广播、连接保持、生命体征上报、SOS、电量提示", "PPG/模拟生命体征、按键、电量", "JSON 上报、OLED 状态、日志"],
            ["救援人员终端", "扫描连接、RSSI 采样、距离等级、趋势显示", "SLE 扫描结果、RSSI、生命体征包", "目标状态、DIST、LEVEL、TREND"],
            ["现场网关", "多终端汇聚、缓存、后台转发、任务下发", "终端数据、后台指令", "MQTT/HTTPS 消息、离线缓存"],
            ["指挥后台", "设备管理、态势展示、告警、优先级、任务派发", "生命体征、位置、环境、任务", "看板、告警、优先级队列、复盘"],
        ],
        [3.0, 5.2, 4.0, 3.8],
    )
    add_heading(doc, "3.2 两板原型功能流程", 2)
    add_paragraph(
        doc,
        "当前原型把两块 WS63 分为受灾端和救援端。受灾端配置为 SLE UART Server，连接建立后每秒生成一条模拟生命体征 JSON；救援端配置为 SLE UART Client，扫描设备名 sle_uart_server，连接后周期读取 RSSI 并接收受灾端数据。救援端根据 RSSI 估算距离分米值，映射为 VERY_NEAR、NEAR、MID、FAR、UNKNOWN 等等级，并根据前后距离差输出 CLOSING、AWAY、STABLE 趋势。",
    )
    add_code_block(
        doc,
        '{"type":"VICTIM","id":"VD-001","seq":12,"spo2":90,"hr":124,"sbp":100,"battery":80,"sos":1}',
    )
    add_heading(doc, "3.3 OLED 界面设计", 2)
    add_paragraph(
        doc,
        "救援端界面优先显示搜救决策字段：连接状态、RSSI、距离、等级、趋势、目标 ID、SpO2、HR、SOS 和电量；受灾端界面优先显示设备身份、连接/广播状态、序号、SpO2、HR、SBP、电量和发送状态。由于 SSD1306 屏幕分辨率有限，界面采用 6x8 字体、分行显示和 1 s 刷新周期，保证信息密度和可读性。",
    )
    add_picture(doc, assets["oled"], "图3-1 救援端与受灾端 OLED 显示布局示意图")
    add_heading(doc, "3.4 通信数据结构", 2)
    add_table_caption(doc, "表3-2 当前原型 JSON 字段说明")
    add_data_table(
        doc,
        ["字段", "含义", "示例", "说明"],
        [
            ["type", "设备/消息类型", "VICTIM", "用于区分受灾端数据"],
            ["id", "受灾终端编号", "VD-001", "救援端 OLED 与日志显示的目标 ID"],
            ["seq", "递增序号", "12", "用于判断数据是否持续更新"],
            ["spo2", "血氧饱和度", "90", "当前为模拟值，工程版接 PPG"],
            ["hr", "心率", "124", "当前为模拟值，单位 bpm"],
            ["sbp", "收缩压趋势", "100", "当前为模拟趋势，不作为医疗诊断"],
            ["battery", "电量百分比", "80", "用于低电量告警"],
            ["sos", "主动求救状态", "1", "1 表示 SOS 有效"],
        ],
        [2.2, 4.0, 2.3, 7.5],
    )
    add_heading(doc, "3.5 优先级与告警功能", 2)
    add_paragraph(
        doc,
        "完整系统后台根据生命风险、静止/失联、位置可达性、环境风险和 SOS 状态计算优先级。评分只是辅助调度，不替代指挥员判断。对于 PPG 估算血压等可信度较低的数据，系统应降低权重并在原因字段中标记数据来源和质量。",
    )
    add_data_table(
        doc,
        ["评分项", "权重", "关键指标", "输出解释"],
        [
            ["生命风险", "45%", "SpO2、HR、SBP/DBP 趋势", "识别缺氧、心率异常、血压风险"],
            ["静止/失联", "20%", "跌倒、长时间静止、数据过期", "判断是否可能失去行动能力"],
            ["位置可达", "15%", "最近救援员距离、定位质量", "辅助选择最合适救援队员"],
            ["环境风险", "10%", "可燃气体、高温、CO、缺氧", "保障救援人员安全"],
            ["SOS", "10%", "主动求救按键", "提高主动求救目标优先级"],
        ],
        [3.0, 2.0, 5.0, 6.0],
    )

    # 4
    add_heading(doc, "4 系统实现", 1)
    add_heading(doc, "4.1 硬件实现", 2)
    add_paragraph(
        doc,
        "基础原型使用两块 WS63 星闪开发板。两端均接入 SSD1306 OLED，I2C 总线为 bus 1，SCL 引脚为 15，SDA 引脚为 16，I2C 速率为 400 kHz。救援端和受灾端均复用项目中的 ssd1306.c 与 ssd1306_fonts.c 驱动文件，实现 128x64 屏幕文字刷新。",
    )
    add_table_caption(doc, "表4-1 硬件与软件技术选型")
    add_data_table(
        doc,
        ["层级", "选型", "作用", "当前状态"],
        [
            ["感知层", "PPG 血氧心率、按键、电量采集", "采集生命体征和 SOS", "原型用模拟值，接口预留"],
            ["传输层", "星闪 NearLink / SLE", "受灾端与救援端短距互联", "已实现 Server/Client 连接"],
            ["显示层", "SSD1306 OLED", "现场显示关键搜救信息", "已实现两端界面"],
            ["控制层", "WS63 LiteOS 工程", "任务调度、通信、显示刷新", "已集成 SLE UART 示例"],
            ["后台层", "Spring Boot + MySQL + Redis + Vue3", "态势、告警、任务和复盘", "完整系统设计"],
        ],
        [2.4, 4.2, 4.3, 5.1],
    )
    add_heading(doc, "4.2 受灾端固件实现", 2)
    add_paragraph(
        doc,
        "受灾端位于 fbb_ws63-shouzaiduan 工程，核心文件为 application/samples/products/sle_uart/sle_uart.c。固件启动后初始化 SLE Server 和 OLED 显示任务，周期任务每 1000 ms 生成一次模拟生命体征数据，并通过 SLE UART 链路发送给已连接的救援端。OLED 同步显示 RESCUE VICTIM、LINK ADVERTISING 或 LINK CONNECTED、ID VD-001、SEQ、SpO2、HR、SBP、BAT 和 SENDING SLE REPORT。",
    )
    add_heading(doc, "4.3 救援端固件实现", 2)
    add_paragraph(
        doc,
        "救援端位于 fbb_ws63-jiuyuanduan 工程，核心文件为 application/samples/products/sle_uart/sle_uart_client/sle_uart_client.c。救援端初始化后启动扫描，发现目标后连接并接收 JSON 数据。解析函数从 payload 中读取 id、spo2、hr、sbp、battery 和 sos，同时读取连接 RSSI 并调用距离估计函数生成 distance、level 和 trend。串口日志以 [RESCUE_TRACK] 为关键字输出，便于调试和验收。",
    )
    add_heading(doc, "4.4 RSSI 粗距离与趋势判断", 2)
    add_paragraph(
        doc,
        "RSSI 估距采用经验模型，输入为连接 RSSI，输出为距离分米值和等级。该方法受人体遮挡、墙体、天线方向和多径影响明显，因此报告将它定义为“粗距离估算”，不写成精确测距。趋势判断比较连续距离估计值：距离明显减小输出 CLOSING，明显增大输出 AWAY，变化较小输出 STABLE。该结果适合引导救援人员“继续靠近或重新搜索”，不适合直接作为地图坐标。",
    )
    add_table_caption(doc, "表4-2 定位能力分层")
    add_data_table(
        doc,
        ["阶段", "数据来源", "输出能力", "适用说明"],
        [
            ["基础版", "SLE 连接 RSSI", "目标发现、粗距离、趋势", "当前两块 WS63 已实现"],
            ["增强版", "SLP/HADM 测距或更可靠距离 API", "更稳定距离值、可提升近场导航", "替换估距数据源即可扩展"],
            ["工程版", "多锚点距离、IMU 轨迹、网关同步", "二维/三维坐标和误差圈", "需要 3 个以上锚点或运动轨迹融合"],
        ],
        [2.3, 4.1, 4.6, 5.0],
    )
    add_heading(doc, "4.5 后台与数据库实现设计", 2)
    add_paragraph(
        doc,
        "完整系统后端采用 Spring Boot，负责设备注册、生命体征入库、定位/测距样本处理、告警去重、优先级计算和任务派发。MySQL 保存现场、人员、设备、生命体征、环境样本、位置样本、告警事件、救援任务和优先级快照；Redis 保存最新状态、活跃告警和优先级队列；前端 Vue3 通过 WebSocket 或 SSE 接收实时事件，在态势总览、受灾人员列表、救援人员面板和任务队列中展示。",
    )
    add_heading(doc, "4.6 测试与验收设计", 2)
    add_table_caption(doc, "表4-3 测试与验收指标")
    add_data_table(
        doc,
        ["测试项", "方法", "期望结果", "判定重点"],
        [
            ["SLE 连接", "先启动受灾端，再启动救援端扫描", "救援端发现 sle_uart_server 并连接", "连接成功率与重连稳定性"],
            ["生命体征上报", "观察串口和 OLED", "每秒更新 JSON、SpO2、HR、SBP、电量", "序号递增且字段解析正确"],
            ["RSSI 趋势", "移动救援端靠近或远离受灾端", "趋势更多出现 CLOSING 或 AWAY", "不要求精确距离，只看趋势合理性"],
            ["OLED 显示", "检查两端屏幕", "不再显示旧温湿度主界面，显示搜救字段", "字段完整、刷新无明显卡顿"],
            ["后台增强", "模拟多终端数据接入", "告警和优先级队列实时刷新", "可解释原因与任务状态一致"],
        ],
        [2.6, 4.4, 4.2, 4.8],
    )

    # 5
    add_heading(doc, "5 其他内容", 1)
    add_heading(doc, "5.1 工业设计与使用方式", 2)
    add_paragraph(
        doc,
        "受灾人员终端可设计为腕带或胸牌形态，外壳需具备防尘、防溅水、抗摔和大按键 SOS 能力；救援人员终端可设计为手持机或臂挂式设备，屏幕优先显示目标状态和报警信息，按键用于目标锁定、确认和任务完成。现场网关采用便携箱式设计，内置电池、星闪模组、4G/Wi-Fi/以太网和本地缓存服务，便于救援队抵达后快速部署。",
    )
    add_heading(doc, "5.2 成本估算", 2)
    add_table_caption(doc, "表5-1 原型与工程样机成本估算")
    add_data_table(
        doc,
        ["部件", "原型方案", "工程增强方案", "说明"],
        [
            ["主控与通信", "WS63 开发板", "定制 MCU + 星闪模组", "竞赛原型优先使用开发板降低调试成本"],
            ["显示", "SSD1306 OLED", "更大 OLED/TFT 或防护屏", "救援端可升级为更强可读性屏幕"],
            ["生命体征", "模拟数据", "MAX30102/AFE 类 PPG + 算法", "血压需明确为趋势或改用专业模块"],
            ["环境检测", "预留接口", "温湿度、工业 LEL、CO/O2", "MQ 类气体模块只能作为风险指数演示"],
            ["供电与结构", "USB/开发板供电", "锂电池、电量计、防护外壳", "决定实际救援可用性"],
        ],
        [3.0, 4.2, 4.5, 4.3],
    )
    add_heading(doc, "5.3 风险与规避措施", 2)
    add_bullets(
        doc,
        [
            "RSSI 受环境影响大：界面显示距离等级和趋势，不显示虚假的高精度坐标；工程版使用 SLP/HADM 或多锚点增强。",
            "生命体征模拟值不能代表真实医学测量：后续接入 PPG 传感器时必须输出信号质量，并在报告和界面中标明数据可信度。",
            "灾后环境可能遮挡严重：通过网关、中继、多锚点和本地缓存提高可用性，后台断网时终端仍可独立显示目标信息。",
            "气体传感器演示风险：低成本模块只显示风险指数，正式应用需工业级 LEL、CO、O2 传感器和校准流程。",
        ],
    )
    add_heading(doc, "5.4 后续开发计划", 2)
    add_paragraph(
        doc,
        "第一阶段完成两块 WS63 的 SLE 连接、模拟生命体征上报、RSSI 粗距离、趋势判断和 OLED 显示；第二阶段接入真实 PPG、IMU、环境传感器和本地报警器，完善救援端目标锁定与扫描提示；第三阶段完成现场网关、后台实时推送、数据库、优先级队列和任务派发；第四阶段加入多锚点定位、断网缓存、外壳、电源管理和现场测试流程，形成可演示、可扩展、可复盘的系统样机。",
    )

    # 6
    add_heading(doc, "6 参考文献", 1)
    refs = [
        "[1] 星闪联盟. 星闪 NearLink 短距无线通信技术白皮书[EB/OL].",
        "[2] 海思技术有限公司. WS63 开发资料与 SLE 示例工程[Z].",
        "[3] 项目组. 基于星闪互联的灾后救援系统开发文档[R]. 2026.",
        "[4] 项目组. 两块 WS63 基础定位版说明[R]. 2026.",
        "[5] GB/T 7714-2015. 信息与文献 参考文献著录规则[S]. 北京: 中国标准出版社, 2015.",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(ref)
        set_run_font(r, size=Pt(10.5), color=DARK)

    props = doc.core_properties
    props.title = "基于星闪互联的灾后救援系统设计报告"
    props.subject = "全国大学生物联网设计竞赛作品报告"
    props.author = "星闪灾后救援系统设计团队"
    props.keywords = "星闪 NearLink; WS63; SLE; 灾后救援; RSSI 粗定位"
    doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    out = build_report()
    print(out)
