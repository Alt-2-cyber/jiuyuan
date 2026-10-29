from pathlib import Path
import importlib.util

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(r"D:\table\jiuyuan")
OUT = ROOT / "项目-按原格式填充.docx"
GENERATOR = ROOT / "tools" / "generate_starflash_report.py"


def load_assets():
    spec = importlib.util.spec_from_file_location("starflash_report_generator", GENERATOR)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.build_assets()


def set_text(paragraph, text):
    if paragraph.runs:
        first = paragraph.runs[0]
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
        first.text = text
    else:
        paragraph.add_run(text)


def clear_para(paragraph):
    set_text(paragraph, "")


def clear_drawings(paragraph):
    for drawing in list(paragraph._p.xpath(".//w:drawing")):
        parent = drawing.getparent()
        if parent is not None:
            parent.remove(drawing)


def strip_numbering(paragraph):
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.numPr is not None:
        ppr.remove(ppr.numPr)


def strip_breaks(paragraph):
    for br in list(paragraph._p.xpath(".//w:br")):
        parent = br.getparent()
        if parent is not None:
            parent.remove(br)


def replace_picture(paragraph, image_path, width_inches=5.6, caption_text=None):
    clear_drawings(paragraph)
    set_text(paragraph, "")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Pt(width_inches * 72))
    if caption_text:
        # If the paragraph originally carried image + caption, keep the caption visible.
        paragraph.add_run("\n" + caption_text)


def remove_paragraph(paragraph):
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_cell(cell, text):
    cell.text = text
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.size = Pt(9)


def fill_table(table, rows):
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            if c < len(table.rows[r].cells):
                set_cell(table.rows[r].cells[c], value)


def remove_range_by_original_index(doc, start, end):
    # Remove by paragraph object snapshot so shifting indices do not matter.
    targets = list(doc.paragraphs[start : end + 1])
    for para in reversed(targets):
        remove_paragraph(para)


def build():
    assets = load_assets()
    src = next(p for p in ROOT.glob("*.docx") if p.name == "项目.docx" or p.stat().st_size == 852099)
    doc = Document(src)
    p = doc.paragraphs
    instruction_paragraphs = list(p[13:20])
    instruction_section_break = p[21]

    # Cover and front matter.
    set_text(p[4], "基于星闪互联的灾后救援系统")
    set_text(p[11], "2026年6月")
    info = doc.tables[0]
    info.cell(0, 1).text = "填写学校名称"
    info.cell(1, 1).text = "星闪灾后救援系统设计团队"
    for r in [3, 4, 5, 6]:
        info.cell(r, 1).text = "填写姓名"

    # Instruction page is removed near the end after all fixed-index edits are done.

    set_text(p[23], "基于星闪互联的灾后救援系统")
    set_text(
        p[26],
        "本作品面向地震、矿难、火灾、坍塌、隧道事故等灾后救援场景，设计一套基于星闪 NearLink 的生命搜救与安全调度系统。系统由受灾人员终端、救援人员终端、现场网关和指挥后台组成，当前原型以两块 WS63 开发板完成最小闭环：受灾端作为 SLE Server 广播并周期发送模拟生命体征 JSON 数据，救援端作为 SLE Client 自动扫描名为 sle_uart_server 的目标、建立连接、读取 RSSI、估算粗距离等级，并根据连续距离变化判断 CLOSING、AWAY、STABLE 趋势；两端均接入 SSD1306 OLED，用于显示连接状态、目标编号、RSSI、距离等级、SpO2、HR、SBP、电量和 SOS 状态。完整系统在此基础上扩展现场网关、Spring Boot 后端、MySQL/Redis 数据存储、Vue3 指挥后台和救援优先级算法，实现本地自治搜救、实时态势展示、告警推送和任务派发。需要说明的是，两块普通 WS63 板当前不能直接输出稳定二维坐标或绝对方向，精确定位需通过 SLP/HADM 测距、多锚点部署或 IMU 轨迹融合进一步增强。",
    )
    set_text(p[27], "关键词：星闪 NearLink；WS63；SLE 通信；RSSI 粗定位；灾后救援；生命体征监测；OLED 显示")

    # TOC placeholders are updated by Word later. Replace visible old labels for readability before update.
    toc_updates = {
        34: "第一章 设计需求分析\t1",
        35: "1.1 灾后救援场景与通信痛点\t1",
        36: "1.2 目标用户与核心需求\t2",
        37: "第二章 特色与创新\t3",
        38: "2.1 星闪近场发现与本地自治\t3",
        39: "2.2 两板基础闭环可平滑升级\t3",
        40: "第三章 功能设计\t5",
        41: "3.1 感知监测与数据上报\t5",
        42: "3.2 搜救显示与预警功能\t5",
        43: "3.3 后台与平台管理功能\t6",
        44: "3.4 典型业务流程\t7",
        45: "第四章 系统实现\t8",
        46: "4.1 整体技术架构实现\t8",
        47: "4.2 硬件系统实现\t8",
        48: "4.3 软件系统实现\t9",
        49: "4.4 数据库与后台增强\t10",
        50: "第五章 其他内容\t11",
        51: "5.1 测试方案与预期指标\t11",
        52: "5.2 成本与产品化分析\t11",
        53: "5.3 风险与规避措施\t12",
        55: "参考文献\t14",
    }
    for idx, text in toc_updates.items():
        set_text(p[idx], text)

    # Chapter 1.
    set_text(p[59], "设计需求分析")
    set_text(p[61], "灾后救援场景与通信痛点")
    set_text(p[62], "传统搜救方式短板")
    set_text(
        p[63],
        "灾后救援现场常存在公网中断、建筑遮挡、烟尘噪声、二次灾害风险高等问题，传统依赖人工呼喊、对讲机或公网定位的搜救方式难以及时发现被困人员并判断救援优先级。人工搜索受环境噪声和视线遮挡影响大；普通短距无线方案缺少目标身份、生命状态和调度闭环；公网或 Wi-Fi 在基站损毁、室内遮挡和复杂电磁环境下可靠性不足。因此，本项目需要在无公网或弱公网条件下，提供本地化目标发现、粗距离判断、生命体征上报和现场可视化提示能力。",
    )
    set_text(p[64], "目标用户与核心需求")
    set_text(
        p[65],
        "本项目面向消防救援、矿山隧道救援、地震坍塌搜救、应急演练和复杂空间巡检等场景，重点满足“发现目标、判断距离、查看生命状态、保障救援人员安全、支持后台调度”的需求。",
    )
    fill_table(
        doc.tables[1],
        [
            ["用户群体", "应用场景", "核心需求"],
            ["消防、矿山、地震救援队", "坍塌、火灾、矿难、隧道事故等灾后现场", "断网环境下快速发现受灾终端，获取距离趋势和生命状态"],
            ["现场救援人员", "进入危险区域搜救被困人员", "手持终端本地显示目标 ID、RSSI、距离等级、趋势、SpO2、HR、SOS"],
            ["救援指挥人员", "多目标、多队伍协同调度", "后台实时态势、告警队列、救援优先级和任务派发"],
            ["受灾或被困人员", "无法主动呼救或通信设备失效", "佩戴终端持续广播身份、电量、生命体征摘要和 SOS 状态"],
        ],
    )
    set_text(p[67], "表 1-1 目标人群与需求分析")
    set_text(
        p[68],
        "结合用户需求，系统需要实现本地自治、低时延短距通信、生命体征上报、距离趋势导航、OLED 可视化和后台增强调度六类能力。当前两板原型优先完成可演示闭环，工程增强版再扩展到多终端、多网关、多锚点和指挥后台。",
    )
    replace_picture(p[69], assets["architecture"], 5.7)
    set_text(p[70], "图1-1 基于星闪互联的灾后救援系统总体架构")
    clear_para(p[72])
    set_text(p[73], "表1-2 当前原型与工程增强能力边界")
    clear_para(p[74])
    fill_table(
        doc.tables[2],
        [
            ["阶段", "核心能力", "硬件/软件基础", "输出结果", "说明"],
            ["基础版", "SLE 连接 + RSSI 粗距离 + 趋势", "两块 WS63、SLE UART、SSD1306 OLED", "发现目标、距离等级、趋势", "当前已实现，适合竞赛演示"],
            ["增强版", "更稳定测距与屏幕导航", "SLP/HADM 测距 API、IMU、目标锁定界面", "更稳定距离、扫描提示、方向趋势", "作为后续升级"],
            ["工程版", "多终端、多网关、多锚点定位和后台调度", "现场网关、Spring Boot、MySQL、Redis、Vue3", "态势总览、告警、任务派发", "需要现场部署和系统联调"],
        ],
    )

    # Chapter 2.
    set_text(p[77], "特色与创新")
    clear_para(p[78])
    set_text(p[79], "星闪近场发现与本地自治")
    set_text(p[80], "SLE 连接与断网可用")
    set_text(
        p[81],
        "本系统以星闪 NearLink/SLE 作为近场通信基础，受灾端作为 SLE Server 广播并等待连接，救援端作为 SLE Client 扫描目标并建立链路。救援人员进入现场后，即使没有公网和后台，也能在 OLED 和串口上看到目标 ID、RSSI、距离等级、趋势和生命体征摘要，实现本地自治搜救。",
    )
    set_text(
        p[82],
        "这种设计把生命安全相关的关键功能前移到终端侧，后台只作为态势汇聚、历史记录、任务派发和优先级分析的增强组件，避免把救援闭环完全绑定在公网或服务器可用性上。",
    )
    set_text(p[83], "两板基础闭环可平滑升级")
    set_text(p[84], "从 RSSI 粗定位到多锚点定位")
    set_text(
        p[85],
        "当前两块 WS63 原型已经完成 SLE 连接、模拟生命体征上报、RSSI 粗距离估计、距离等级和趋势判断。RSSI 估距受墙体、人体遮挡、天线方向和多径影响明显，因此报告中明确定位为“粗距离和趋势辅助”，不宣称稳定二维坐标或绝对方向。后续若接入 SLP/HADM 测距 API，可直接替换距离来源；若部署 3 个以上锚点或融合 IMU 运动轨迹，可进一步实现二维定位和误差圈显示。",
    )
    set_text(p[86], "生命体征与救援优先级创新")
    set_text(
        p[87],
        "系统不是单纯的无线测距 Demo，而是把目标身份、SpO2、HR、SBP 趋势、电量、SOS、RSSI、距离等级和趋势放在同一条链路中处理。后台增强版根据生命风险、静止/失联、位置可达性、环境风险和 SOS 形成可解释救援优先级，帮助指挥员判断先救谁、派谁去、是否需要安全防护。血压字段在当前原型中为模拟/趋势数据，不作为医疗诊断结论。",
    )
    set_text(p[88], "表 2-1 作品核心创新点")
    fill_table(
        doc.tables[3],
        [
            ["创新点", "实现方式", "价值"],
            ["星闪 SLE 近场发现", "受灾端广播，救援端扫描并连接 sle_uart_server", "断网场景下仍可发现目标并获取状态"],
            ["RSSI 粗距离与趋势导航", "读取连接 RSSI，输出 CLOSING/AWAY/STABLE", "帮助判断移动方向是否更接近目标"],
            ["生命体征与定位信息同链路上报", "JSON 携带 id、spo2、hr、sbp、battery、sos", "同时判断身份、生命状态和紧急程度"],
            ["双端 OLED 本地显示", "SSD1306 显示 RESCUE TRACK / RESCUE VICTIM 页面", "现场无需电脑即可查看关键搜救信息"],
            ["可解释救援优先级", "按生命风险、位置可达性、环境风险、SOS 加权评分", "后台增强版支持透明调度和人工复核"],
        ],
    )
    replace_picture(p[92], assets["workflow"], 5.6, "图2-1 两块 WS63 基础定位工作流程")

    # Chapter 3.
    set_text(p[93], "功能设计")
    clear_para(p[95])
    set_text(p[96], "感知监测与数据上报")
    set_text(p[97], "受灾端数据采集与模拟上报")
    set_text(p[98], "受灾人员终端承担身份广播、生命体征摘要上报、SOS 状态提示和电量状态提示功能。当前原型使用模拟数据完成闭环，后续可接入 PPG 血氧心率传感器、SOS 按键、电池电量计和 IMU。")
    set_text(p[99], "生命体征上报：当前 JSON 字段包含 id、seq、spo2、hr、sbp、battery、sos，救援端每秒接收并解析。")
    set_text(p[100], "连接状态显示：受灾端 OLED 显示 RESCUE VICTIM、LINK ADVERTISING 或 LINK CONNECTED、ID、SEQ、SpO2、HR、SBP、BAT。")
    set_text(p[101], "数据边界：SpO2、HR、SBP 当前为演示模拟值；SBP 应描述为血压趋势辅助，不作为医疗诊断。")
    set_text(p[102], "扩展接口：后续可加入真实 PPG、IMU、环境或皮温传感器，并在协议中加入 signalQuality 字段。")
    set_text(p[103], "搜救显示与预警功能")
    set_text(p[104], "救援端目标发现与趋势提示")
    set_text(p[105], "救援端自动扫描目标，发现 sle_uart_server 后建立连接并读取 RSSI。")
    set_text(p[106], "RSSI 估算：将 RSSI 映射为 VERY_NEAR、NEAR、MID、FAR、UNKNOWN 等距离等级。")
    set_text(p[107], "趋势判断：连续距离估计值下降显示 CLOSING，上升显示 AWAY，变化较小显示 STABLE。")
    set_text(p[108], "OLED 显示：救援端显示 RESCUE TRACK、连接状态、RSSI、DIST、LEVEL、TREND、ID、SpO2、HR、SOS/BAT。")
    set_text(p[109], "后台与平台管理功能")
    set_text(p[110], "现场网关与指挥后台规划")
    set_text(p[111], "现场网关：汇聚受灾终端、救援终端和环境数据，在公网不稳定时进行本地缓存，恢复链路后通过 MQTT/HTTPS 上传后台。")
    set_text(p[112], "指挥后台：采用 Spring Boot、MySQL、Redis、Vue3 等技术，提供设备管理、实时态势、告警队列、救援优先级、任务派发和复盘导出。")

    # Chapter 4.
    set_text(p[118], "系统实现")
    clear_para(p[119])
    set_text(p[120], "整体技术架构实现")
    set_text(p[121], "物联网分层架构实现")
    set_text(p[122], "系统采用感知层、传输层、控制层、应用层四层架构。感知层由 WS63 受灾端、WS63 救援端和 OLED 显示组成；传输层以星闪 SLE 连接为核心；控制层完成 JSON 编解码、RSSI 估距、趋势判断和显示刷新；应用层由现场网关和指挥后台完成态势展示、告警和任务派发。")
    replace_picture(p[124], assets["architecture"], 5.6)
    set_text(p[125], "图 4-1 系统整体技术架构图")
    set_text(p[126], "硬件系统实现")
    set_text(p[127], "WS63 双板与 OLED 接线")
    set_text(p[128], "硬件由主控通信、显示交互、生命体征输入和扩展传感四部分组成：")
    set_text(p[129], "核心主控：两块 WS63 星闪开发板，受灾端配置为 SLE Server，救援端配置为 SLE Client。")
    set_text(p[130], "显示模块：两端均接入 SSD1306 OLED，I2C bus 1，SCL 引脚 15，SDA 引脚 16，速率 400 kHz。")
    set_text(p[131], "生命体征：当前使用模拟 SpO2、HR、SBP、电量和 SOS，后续接入 PPG 与按键。")
    set_text(p[132], "通信模块：使用 SLE UART 示例完成扫描、连接、收发和 RSSI 读取，救援端根据 RSSI 输出距离等级与趋势。")
    set_text(p[133], "当前屏幕已经从温湿度演示界面改为救援信息界面，受灾端显示生命体征上报状态，救援端显示目标追踪信息。")
    replace_picture(p[136], assets["oled"], 5.5)
    set_text(p[137], "图 4-2 救援端与受灾端 OLED 显示界面")
    set_text(p[138], "软件系统实现")
    set_text(p[139], "固件、协议与后台增强")
    set_text(p[140], "受灾端固件：位于 fbb_ws63-shouzaiduan 工程，核心文件 application/samples/products/sle_uart/sle_uart.c。固件启动 OLED 显示任务和周期上报任务，每 1000 ms 发送一条模拟生命体征 JSON。")
    set_text(p[141], "救援端固件：位于 fbb_ws63-jiuyuanduan 工程，核心文件 application/samples/products/sle_uart/sle_uart_client/sle_uart_client.c。救援端扫描目标、连接后接收 payload、解析生命体征、读取 RSSI 并打印 [RESCUE_TRACK] 日志。")
    replace_picture(p[142], assets["workflow"], 5.5)
    set_text(p[143], "图 4-3 两板基础定位与数据上报流程")
    set_text(p[144], "后台增强：现场网关通过 MQTT/HTTPS 上传设备最新状态，Spring Boot 后端写入 MySQL 与 Redis，Vue3 指挥后台通过 WebSocket/SSE 获取实时事件。")
    replace_picture(p[146], assets["priority"], 5.6)
    set_text(p[147], "图 4-4 救援优先级计算结构")
    set_text(p[149], "数据库设计包括 site、person、device、vital_sample、env_sample、location_sample、alert_event、rescue_task 和 priority_snapshot 等表。优先级评分综合生命风险、静止/失联、位置可达性、环境风险和 SOS，并输出可解释原因。")

    # Remove the long blank gap after chapter 4, then refetch paragraph list.
    remove_range_by_original_index(doc, 151, 175)
    p = doc.paragraphs

    # The old template tail starts at original p176, now at current index 151.
    set_text(p[151], "其他内容")
    clear_para(p[153])
    set_text(p[154], "测试方案与预期指标")
    clear_para(p[155])
    clear_para(p[156])
    set_text(p[157], "原型测试指标")
    set_text(
        p[158],
        "测试重点包括：SLE 扫描连接成功率、生命体征 JSON 是否每秒更新、RSSI 距离等级是否随移动变化、趋势 CLOSING/AWAY/STABLE 是否符合靠近远离过程、OLED 是否稳定显示 RESCUE TRACK 与 RESCUE VICTIM 页面、串口日志是否包含 [RESCUE_TRACK] 和 [RESCUE_VICTIM] 关键字。验收时不以绝对距离精度作为唯一指标，而以发现目标、持续上报、趋势合理和本地可视化为主要判断标准。",
    )
    clear_para(p[159])
    clear_para(p[160])
    p[161].style = "Heading 2"
    p[161].paragraph_format.page_break_before = True
    set_text(p[161], "成本与产品化分析")
    p[162].style = "Normal"
    strip_numbering(p[162])
    set_text(
        p[162],
        "产品化阶段需将开发板方案替换为定制主控与星闪模组，加入防护外壳、锂电池、电量计、SOS 按键、真实 PPG 传感器和工业级环境传感器。RSSI 估距必须标注为粗略参考；若用于正式定位，应升级为 SLP/HADM、多锚点或 IMU 轨迹融合。生命体征和血压趋势不能替代医疗诊断，气体检测在工程应用中必须采用工业级传感器并建立校准流程。",
    )

    p[163].style = "Heading 1"
    set_text(p[163], "参考文献")
    for ref_idx in range(164, 169):
        p[ref_idx].style = "Normal"
        strip_numbering(p[ref_idx])
    set_text(p[164], "[1] 星闪联盟. 星闪 NearLink 短距无线通信技术白皮书[EB/OL].")
    set_text(p[165], "[2] 海思技术有限公司. WS63 开发资料与 SLE 示例工程[Z].")
    set_text(p[166], "[3] 项目组. 基于星闪互联的灾后救援系统开发文档[R]. 2026.")
    set_text(p[167], "[4] 项目组. 两块 WS63 基础定位版说明[R]. 2026.")
    set_text(p[168], "[5] GB/T 7714-2015. 信息与文献 参考文献著录规则[S]. 北京: 中国标准出版社, 2015.")
    for idx in range(len(doc.paragraphs) - 1, 168, -1):
        remove_paragraph(doc.paragraphs[idx])

    # Delete the visible "填写说明" page while preserving the cover page break.
    strip_breaks(instruction_section_break)
    clear_para(instruction_section_break)
    for para in reversed(instruction_paragraphs):
        remove_paragraph(para)

    for para in doc.paragraphs:
        if not para.style.name.startswith("Heading") and para.style.name not in ("toc 1", "toc 2", "toc 3"):
            strip_numbering(para)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
